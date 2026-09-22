"""
Importación masiva de estudiantes y notas desde archivos Excel (.xlsx) o CSV.

Formato esperado: una fila por (estudiante, curso).  Columnas admitidas
(los nombres no distinguen mayúsculas ni acentos):

    grado, anio, apellidos, nombres, codigo, curso, u1, u2, u3, u4, observaciones

- 'grado' y 'anio' pueden omitirse si la importación se hace desde un grado concreto.
- Si un estudiante tiene varios cursos, se repite en varias filas (mismos datos
  de estudiante, distinto 'curso').
- Los cursos pertenecen al GRADO: si un curso del archivo no existe en el grado se
  agrega a su catálogo y queda disponible para todos sus estudiantes.
- Las notas vacías o no numéricas se dejan sin registrar (None). Escala 0-100.
- 'observaciones' se aplica al estudiante (gana el último valor no vacío).
"""

import csv
import io
import unicodedata
from datetime import datetime

import db

COLUMNAS_PLANTILLA = ["grado", "anio", "apellidos", "nombres", "codigo",
                      "curso", "u1", "u2", "u3", "u4", "observaciones"]

_SINONIMOS = {
    "grado": "grado", "grados": "grado", "seccion": "grado", "seccion/grado": "grado",
    "anio": "anio", "ano": "anio", "year": "anio", "ciclo": "anio", "cicloescolar": "anio",
    "apellidos": "apellidos", "apellido": "apellidos",
    "nombres": "nombres", "nombre": "nombres",
    "codigo": "codigo", "carnet": "codigo", "id": "codigo", "codigoestudiante": "codigo",
    "curso": "curso", "materia": "curso", "asignatura": "curso", "clase": "curso",
    "u1": "u1", "unidad1": "u1", "i": "u1", "primeraunidad": "u1", "iunidad": "u1",
    "u2": "u2", "unidad2": "u2", "ii": "u2", "segundaunidad": "u2", "iiunidad": "u2",
    "u3": "u3", "unidad3": "u3", "iii": "u3", "terceraunidad": "u3", "iiiunidad": "u3",
    "u4": "u4", "unidad4": "u4", "iv": "u4", "cuartaunidad": "u4", "ivunidad": "u4",
    "observaciones": "observaciones", "observacion": "observaciones",
    "comportamiento": "observaciones", "observacionesdecomportamiento": "observaciones",
    "observacionescomportamiento": "observaciones",
}


def _norm(texto):
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return "".join(ch for ch in t.lower() if ch.isalnum())


def _key(texto):
    """Clave de comparación tolerante a acentos, mayúsculas y espacios."""
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(t.lower().split())


def _mapear_encabezados(fila_encabezados):
    mapa = {}
    for idx, col in enumerate(fila_encabezados):
        clave = _SINONIMOS.get(_norm(col))
        if clave and clave not in mapa:
            mapa[clave] = idx
    return mapa


def parse_nota(valor):
    """Devuelve float 0-100, None si vacío, o 'ERROR' si no es interpretable."""
    if valor is None:
        return None
    s = str(valor).strip().replace(",", ".")
    if s == "" or s.lower() in ("np", "n/a", "na", "-", "--", "s/n"):
        return None
    try:
        n = float(s)
    except ValueError:
        return "ERROR"
    if n < 0 or n > 100:
        return "ERROR"
    return round(n, 2)


def _leer_filas_crudas(nombre_archivo, contenido_bytes):
    """Devuelve lista de listas (incluida la fila de encabezados)."""
    nombre = (nombre_archivo or "").lower()
    if nombre.endswith(".xlsx") or nombre.endswith(".xlsm"):
        try:
            import openpyxl
        except ImportError:
            raise ValueError("Para leer archivos Excel instale 'openpyxl' "
                             "(pip install openpyxl). También puede usar CSV.")
        wb = openpyxl.load_workbook(io.BytesIO(contenido_bytes), read_only=True, data_only=True)
        ws = wb.active
        filas = []
        for fila in ws.iter_rows(values_only=True):
            filas.append(["" if c is None else c for c in fila])
        wb.close()
        return filas
    # CSV / texto
    texto = None
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            texto = contenido_bytes.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if texto is None:
        raise ValueError("No se pudo leer el archivo (codificación desconocida).")
    try:
        dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=";,\t|")
        delim = dialecto.delimiter
    except csv.Error:
        delim = ";" if texto.count(";") >= texto.count(",") else ","
    return [list(r) for r in csv.reader(io.StringIO(texto), delimiter=delim)]


def parse_archivo(nombre_archivo, contenido_bytes, forzar_grado=None, forzar_anio=None):
    """
    Devuelve (filas, errores_formato).
    filas: lista de dicts con claves grado, anio, apellidos, nombres, codigo,
           curso, u1..u4 (float|None), observaciones, _fila (nº de fila en el archivo).
    errores_formato: lista de strings (problemas que impiden procesar filas concretas).
    """
    crudas = _leer_filas_crudas(nombre_archivo, contenido_bytes)
    crudas = [f for f in crudas if any(str(c).strip() for c in f)]
    if not crudas:
        raise ValueError("El archivo está vacío.")

    mapa = _mapear_encabezados(crudas[0])
    obligatorias = ["apellidos", "nombres"]
    if not forzar_grado:
        obligatorias = ["grado"] + obligatorias
    faltan = [c for c in obligatorias if c not in mapa]
    if faltan:
        raise ValueError("Faltan columnas obligatorias en el archivo: " + ", ".join(faltan) +
                         ". Descargue la plantilla de ejemplo.")

    filas, errores = [], []
    for i, cruda in enumerate(crudas[1:], start=2):
        def val(clave):
            idx = mapa.get(clave)
            if idx is None or idx >= len(cruda):
                return ""
            return str(cruda[idx]).strip()

        apellidos = val("apellidos")
        nombres = val("nombres")
        if not apellidos and not nombres:
            continue
        if not apellidos or not nombres:
            errores.append(f"Fila {i}: falta apellidos o nombres.")
            continue

        grado = forzar_grado or val("grado")
        anio_txt = str(forzar_anio) if forzar_anio else val("anio")
        try:
            anio = int(float(anio_txt)) if anio_txt else datetime.now().year
        except ValueError:
            errores.append(f"Fila {i}: año inválido «{anio_txt}».")
            continue
        if not grado:
            errores.append(f"Fila {i}: falta el grado.")
            continue

        notas = {}
        for u in ("u1", "u2", "u3", "u4"):
            n = parse_nota(val(u)) if u in mapa else None
            if n == "ERROR":
                errores.append(f"Fila {i}: nota {u.upper()} «{val(u)}» no es un número entre "
                               f"0 y 100; se dejó sin registrar.")
                n = None
            notas[u] = n

        observaciones = val("observaciones")
        if len(observaciones.split()) > 500:
            errores.append(f"Fila {i}: las observaciones superan las 500 palabras.")
            continue

        filas.append({
            "grado": grado, "anio": anio,
            "apellidos": apellidos, "nombres": nombres,
            "codigo": val("codigo"),
            "curso": val("curso"),
            "u1": notas["u1"], "u2": notas["u2"], "u3": notas["u3"], "u4": notas["u4"],
            "observaciones": observaciones,
            "_fila": i,
        })
    return filas, errores


def aplicar_import(conn, usuario, filas, dry_run=True):
    """Aplica (o simula) la importación. Devuelve un dict-resumen."""
    ahora = datetime.now().isoformat(timespec="seconds")
    res = {
        "dry_run": dry_run, "total_filas": len(filas),
        "grados_creados": [], "grados_rechazados": [],
        "estudiantes_creados": 0, "estudiantes_actualizados": 0,
        "cursos_creados": 0, "cursos_actualizados": 0, "notas_asignadas": 0,
        "errores": [],
    }
    # cursos_creados      -> cursos agregados al catálogo del grado
    # cursos_actualizados -> filas de notas (estudiante-curso) modificadas
    # notas_asignadas     -> filas de notas con al menos una unidad con valor
    es_admin = usuario["rol"] == "admin"

    # ---- Resolver grados (comparación tolerante a acentos y mayúsculas) ----
    grados_idx = {}
    for g in conn.execute("SELECT id, nombre, anio FROM grados WHERE activo = 1").fetchall():
        grados_idx[(_key(g["nombre"]), g["anio"])] = g["id"]

    grado_de = {}   # (_key(nombre), anio) -> gid real  |  -1 (se creará)  |  None (rechazado)
    for nombre, anio in {(f["grado"].strip(), f["anio"]) for f in filas}:
        k = (_key(nombre), anio)
        if k in grado_de:
            continue
        gid = grados_idx.get(k)
        if gid:
            if not es_admin:
                acc = conn.execute(
                    "SELECT nivel FROM accesos_grado WHERE usuario_id = ? AND grado_id = ? AND activo = 1",
                    (usuario["id"], gid),
                ).fetchone()
                if not acc or acc["nivel"] != "admin":
                    res["grados_rechazados"].append(
                        f"{nombre} ({anio}): no tiene acceso de edición a este grado.")
                    grado_de[k] = None
                    continue
            grado_de[k] = gid
        elif dry_run:
            grado_de[k] = -1
            res["grados_creados"].append(f"{nombre} ({anio})")
        else:
            cur = conn.execute(
                "INSERT INTO grados (nombre, anio, activo, creado_por, creado_en) VALUES (?, ?, 1, ?, ?)",
                (nombre.strip(), anio, usuario["id"], ahora),
            )
            gid = cur.lastrowid
            if not es_admin:
                conn.execute(
                    "INSERT INTO accesos_grado (usuario_id, grado_id, nivel, activo, creado_en) "
                    "VALUES (?, ?, 'admin', 1, ?)", (usuario["id"], gid, ahora),
                )
            grados_idx[k] = gid
            grado_de[k] = gid
            res["grados_creados"].append(f"{nombre} ({anio})")

    # ---- Índices por grado (cargados una sola vez) ----
    _cursos_idx, _est_idx = {}, {}

    def cursos_de(gid):
        if gid not in _cursos_idx:
            d = {}
            for cg in conn.execute(
                "SELECT id, nombre, activo FROM cursos_grado WHERE grado_id = ?", (gid,)
            ).fetchall():
                d[_key(cg["nombre"])] = {"id": cg["id"], "activo": cg["activo"]}
            _cursos_idx[gid] = d
        return _cursos_idx[gid]

    def estudiantes_de(gid):
        if gid not in _est_idx:
            d = {}
            for e in conn.execute(
                "SELECT id, nombres, apellidos, codigo FROM estudiantes WHERE grado_id = ? AND activo = 1", (gid,)
            ).fetchall():
                if e["codigo"]:
                    d["c|" + _key(e["codigo"])] = e["id"]
                d["n|" + _key(e["apellidos"]) + "|" + _key(e["nombres"])] = e["id"]
            _est_idx[gid] = d
        return _est_idx[gid]

    sim_cursos = {}   # k_grado -> set(_key curso)  (para grados que se crearán)
    sim_est = {}      # k_grado -> set(clave estudiante)
    procesado_est = {}  # (gid, clave) -> eid resuelto, para no recontar en filas siguientes
    afectados = set()

    for f in filas:
        kg = (_key(f["grado"].strip()), f["anio"])
        gid = grado_de.get(kg)
        if gid is None:
            continue
        real = isinstance(gid, int) and gid > 0

        s_cod = "c|" + _key(f["codigo"]) if f["codigo"] else None
        s_nom = "n|" + _key(f["apellidos"]) + "|" + _key(f["nombres"])
        clave_est = s_cod or s_nom

        # -- Estudiante --
        eid = procesado_est.get((gid, s_cod)) or procesado_est.get((gid, s_nom))
        if eid is None:
            if real:
                idx = estudiantes_de(gid)
                encontrado = (idx.get(s_cod) if s_cod else None) or idx.get(s_nom)
                if isinstance(encontrado, int):
                    eid = encontrado
                    res["estudiantes_actualizados"] += 1
                    if not dry_run:
                        if f["codigo"]:
                            conn.execute("UPDATE estudiantes SET codigo = ? WHERE id = ?", (f["codigo"], eid))
                        if f["observaciones"]:
                            conn.execute(
                                "UPDATE estudiantes SET observaciones_comportamiento = ? WHERE id = ?",
                                (f["observaciones"], eid))
                else:
                    res["estudiantes_creados"] += 1
                    if dry_run:
                        eid = "SIM"
                        if s_cod:
                            idx[s_cod] = "SIM"
                        idx[s_nom] = "SIM"
                    else:
                        cur = conn.execute(
                            """INSERT INTO estudiantes (grado_id, nombres, apellidos, codigo,
                                   observaciones_comportamiento, activo, creado_en)
                               VALUES (?, ?, ?, ?, ?, 1, ?)""",
                            (gid, f["nombres"], f["apellidos"], f["codigo"],
                             f["observaciones"] or None, ahora))
                        eid = cur.lastrowid
                        if s_cod:
                            idx[s_cod] = eid
                        idx[s_nom] = eid
            else:  # grado que se creará: solo simular
                sset = sim_est.setdefault(kg, set())
                if clave_est not in sset:
                    sset.add(clave_est)
                    res["estudiantes_creados"] += 1
                eid = "SIM"
            procesado_est[(gid, s_nom)] = eid
            if s_cod:
                procesado_est[(gid, s_cod)] = eid
        elif not dry_run and isinstance(eid, int) and f["observaciones"]:
            conn.execute("UPDATE estudiantes SET observaciones_comportamiento = ? WHERE id = ?",
                         (f["observaciones"], eid))

        # -- Curso del grado --
        if not f["curso"]:
            continue
        curso = f["curso"].strip()
        ckey = _key(curso)
        cgid = None
        if real:
            cursos = cursos_de(gid)
            info = cursos.get(ckey)
            if info:
                cgid = info["id"]
                if not dry_run and info["id"] and not info["activo"]:
                    conn.execute("UPDATE cursos_grado SET activo = 1 WHERE id = ?", (cgid,))
                    info["activo"] = 1
            else:
                res["cursos_creados"] += 1
                if dry_run:
                    cursos[ckey] = {"id": None, "activo": 1}
                else:
                    maxorden = conn.execute(
                        "SELECT COALESCE(MAX(orden), 0) m FROM cursos_grado WHERE grado_id = ?", (gid,)
                    ).fetchone()["m"]
                    cur = conn.execute(
                        "INSERT INTO cursos_grado (grado_id, nombre, orden, activo) VALUES (?, ?, ?, 1)",
                        (gid, curso, maxorden + 1))
                    cgid = cur.lastrowid
                    cursos[ckey] = {"id": cgid, "activo": 1}
        else:
            sset = sim_cursos.setdefault(kg, set())
            if ckey not in sset:
                sset.add(ckey)
                res["cursos_creados"] += 1

        if any(f[u] is not None for u in ("u1", "u2", "u3", "u4")):
            res["notas_asignadas"] += 1

        if not dry_run and real and isinstance(eid, int) and isinstance(cgid, int):
            afectados.add(gid)
            nrow = conn.execute(
                "SELECT id FROM notas WHERE estudiante_id = ? AND curso_grado_id = ?", (eid, cgid)
            ).fetchone()
            if nrow:
                res["cursos_actualizados"] += 1
                conn.execute("UPDATE notas SET u1=?, u2=?, u3=?, u4=? WHERE id=?",
                             (f["u1"], f["u2"], f["u3"], f["u4"], nrow["id"]))
            else:
                conn.execute(
                    "INSERT INTO notas (estudiante_id, curso_grado_id, u1, u2, u3, u4) VALUES (?, ?, ?, ?, ?, ?)",
                    (eid, cgid, f["u1"], f["u2"], f["u3"], f["u4"]))

    # Completar filas de notas faltantes (otros estudiantes / otros cursos del grado)
    if not dry_run:
        for gid in afectados:
            db.backfill_notas(conn, gid)

    return res


def plantilla_csv():
    out = io.StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow(COLUMNAS_PLANTILLA)
    w.writerow(["3 Primaria A", datetime.now().year, "Bianchi", "Luca", "001",
                "Matemática", 85, 90, 78, 92, "Excelente participación."])
    w.writerow(["3 Primaria A", datetime.now().year, "Bianchi", "Luca", "001",
                "Lengua", 70, 65, 80, 75, ""])
    w.writerow(["3 Primaria A", datetime.now().year, "Rossi", "Giulia", "002",
                "Matemática", 55, 60, 58, 62, "Debe reforzar tareas en casa."])
    return out.getvalue().encode("utf-8-sig")


def plantilla_xlsx():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Notas"
    ws.append(COLUMNAS_PLANTILLA)
    ws.append(["3 Primaria A", datetime.now().year, "Bianchi", "Luca", "001",
               "Matemática", 85, 90, 78, 92, "Excelente participación."])
    ws.append(["3 Primaria A", datetime.now().year, "Bianchi", "Luca", "001",
               "Lengua", 70, 65, 80, 75, ""])
    ws.append(["3 Primaria A", datetime.now().year, "Rossi", "Giulia", "002",
               "Matemática", 55, 60, 58, 62, "Debe reforzar tareas en casa."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
